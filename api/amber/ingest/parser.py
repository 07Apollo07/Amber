"""Document parsers for PDF, Word (.docx), PowerPoint (.pptx), Excel (.xlsx), TXT, and Markdown files."""

import functools
import io
import os
import re
import zipfile
from typing import Any

from docx import Document as DocxDocument
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph as DocxParagraph
from openpyxl import load_workbook
from pypdf import PdfReader

DOCX_MIME = "wordprocessingml"
PPTX_MIME = "presentationml"
XLSX_MIME = "spreadsheetml"
SLIDE_MARKER = re.compile(r"<!-- Slide number: (\d+) -->")


def parse_document(file_bytes: bytes, mime_type: str, filename: str) -> list[dict[str, Any]]:
    """Parses file bytes into a list of section/page blocks with content and metadata."""
    file_type = _infer_file_type(mime_type, filename)

    if file_type == "pdf":
        return _parse_pdf(file_bytes)
    elif file_type == "docx":
        return _parse_docx(file_bytes, filename)
    elif file_type == "pptx":
        return _parse_pptx(file_bytes, filename)
    elif file_type == "xlsx":
        return _parse_xlsx(file_bytes, filename)
    elif file_type in ("txt", "md"):
        text = file_bytes.decode("utf-8", errors="ignore")
        return [{"page": 1, "section_title": filename, "content": text}]
    else:
        text = file_bytes.decode("utf-8", errors="ignore")
        return [{"page": 1, "section_title": filename, "content": text}]


def _infer_file_type(mime_type: str, filename: str) -> str:
    lower_name = filename.lower()
    if "pdf" in mime_type or lower_name.endswith(".pdf"):
        return "pdf"
    if DOCX_MIME in mime_type or lower_name.endswith(".docx"):
        return "docx"
    if PPTX_MIME in mime_type or lower_name.endswith(".pptx"):
        return "pptx"
    if XLSX_MIME in mime_type or lower_name.endswith(".xlsx"):
        return "xlsx"
    if "markdown" in mime_type or lower_name.endswith(".md"):
        return "md"
    if "text" in mime_type or lower_name.endswith(".txt"):
        return "txt"
    return "other"


def _parse_pdf(file_bytes: bytes) -> list[dict[str, Any]]:
    reader = PdfReader(io.BytesIO(file_bytes))
    blocks = []
    for idx, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        if text.strip():
            blocks.append(
                {
                    "page": idx + 1,
                    "section_title": f"Page {idx + 1}",
                    "content": text,
                }
            )
    return blocks


def _parse_docx(file_bytes: bytes, filename: str) -> list[dict[str, Any]]:
    """One block per heading section, in reading order, with tables flattened to "cell | cell" rows.
    Word files have no fixed pages, so blocks carry no page number."""
    try:
        doc = DocxDocument(io.BytesIO(file_bytes))
    except Exception as exc:
        raise ValueError(f"Couldn't read {filename} as a Word document: {exc}") from exc

    blocks: list[dict[str, Any]] = []
    title = filename
    lines: list[str] = []

    def flush() -> None:
        content = "\n".join(lines).strip()
        if content:
            blocks.append({"page": None, "section_title": title, "content": content})
        lines.clear()

    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            paragraph = DocxParagraph(child, doc)
            text = paragraph.text.strip()
            if not text:
                continue
            style = (paragraph.style.name if paragraph.style is not None else "") or ""
            if style.startswith(("Heading", "Title")):
                flush()
                title = text
            lines.append(text)
        elif tag == "tbl":
            for row in DocxTable(child, doc).rows:
                cells = list(dict.fromkeys(c.text.strip() for c in row.cells if c.text.strip()))  # merged cells repeat
                if cells:
                    lines.append(" | ".join(cells))
    flush()
    return blocks


@functools.cache
def _markitdown() -> Any:
    """Imported on first use. markitdown's dependencies load api/.env into os.environ on import, which would
    bypass the app's own settings rules (for example, prod ignores dotenv), so undo any such changes."""
    before = dict(os.environ)
    try:
        from markitdown import MarkItDown

        return MarkItDown(enable_plugins=False)
    finally:
        os.environ.clear()
        os.environ.update(before)


def _parse_pptx(file_bytes: bytes, filename: str) -> list[dict[str, Any]]:
    """Converts the deck to Markdown with MarkItDown, then splits it into one block per slide
    (page = slide number). MarkItDown includes tables and speaker notes."""
    if not zipfile.is_zipfile(io.BytesIO(file_bytes)):  # MarkItDown would fall back to plain text
        raise ValueError(f"Couldn't read {filename} as a PowerPoint presentation: not a .pptx file")
    try:
        markdown = _markitdown().convert_stream(io.BytesIO(file_bytes), file_extension=".pptx").markdown
    except Exception as exc:
        raise ValueError(f"Couldn't read {filename} as a PowerPoint presentation: {exc}") from exc

    # re.split with one group gives [preamble, "1", slide 1 text, "2", slide 2 text, ...]
    parts = SLIDE_MARKER.split(markdown)
    blocks = []
    for number, text in zip(parts[1::2], parts[2::2]):
        content = text.strip()
        if not content:
            continue
        heading = re.match(r"#\s+(.+)", content)
        blocks.append(
            {
                "page": int(number),
                "section_title": f"Slide {number}: {heading.group(1).strip()}" if heading else f"Slide {number}",
                "content": content,
            }
        )
    return blocks


def _parse_xlsx(file_bytes: bytes, filename: str) -> list[dict[str, Any]]:
    """One block per sheet (page = sheet position), rows as "cell | cell" lines. Formulas read as their
    last saved values; empty rows and fully empty sheets are skipped."""
    try:
        workbook = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError(f"Couldn't read {filename} as an Excel workbook: {exc}") from exc

    blocks = []
    try:
        for idx, sheet in enumerate(workbook.worksheets, start=1):
            lines = []
            for row in sheet.iter_rows(values_only=True):
                cells = [str(v).strip() for v in row if v is not None and str(v).strip()]
                if cells:
                    lines.append(" | ".join(cells))
            if lines:
                blocks.append({"page": idx, "section_title": f"Sheet: {sheet.title}", "content": "\n".join(lines)})
    finally:
        workbook.close()
    return blocks
