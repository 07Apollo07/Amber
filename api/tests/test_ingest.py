"""Tests for document parsing, chunking, and ingestion API endpoints."""

from fastapi.testclient import TestClient

from amber.ingest.chunker import chunk_document_blocks
from amber.ingest.parser import parse_document


def test_parse_text_document():
    content = b"Hello world! This is a test markdown document.\n\n# Heading 1\nSome paragraph text here."
    blocks = parse_document(content, "text/markdown", "test.md")
    assert len(blocks) > 0
    assert any("Hello world" in b["content"] for b in blocks)


def test_chunk_document_blocks():
    blocks = [{"content": "Paragraph one. " * 50, "section_title": "Section 1", "metadata": {}}]
    chunks = chunk_document_blocks(blocks)
    assert len(chunks) > 0
    assert all("content" in c and "token_count" in c for c in chunks)


def test_documents_endpoints_unauthenticated(client: TestClient):
    response = client.get("/api/documents")
    assert response.status_code == 401


def _docx_bytes() -> bytes:
    import io

    from docx import Document

    doc = Document()
    doc.add_heading("Quarterly Plan", level=1)
    doc.add_paragraph("Revenue grew by twelve percent.")
    doc.add_heading("Risks", level=2)
    doc.add_paragraph("Supplier delays.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Region", "Sales"
    table.cell(1, 0).text, table.cell(1, 1).text = "EMEA", "42"
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def _pptx_bytes() -> bytes:
    import io

    from pptx import Presentation

    prs = Presentation()
    first = prs.slides.add_slide(prs.slide_layouts[1])
    first.shapes.title.text = "Roadmap"
    first.placeholders[1].text = "Ship parsing support"
    first.notes_slide.notes_text_frame.text = "Mention the pilot customers"
    prs.slides.add_slide(prs.slide_layouts[6])  # blank slide is skipped
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def test_parse_docx_sections_and_tables():
    blocks = parse_document(_docx_bytes(), "application/octet-stream", "plan.docx")
    titles = [b["section_title"] for b in blocks]
    assert titles == ["Quarterly Plan", "Risks"]
    assert "Revenue grew" in blocks[0]["content"]
    assert "EMEA | 42" in blocks[1]["content"] and "Region | Sales" in blocks[1]["content"]
    assert chunk_document_blocks(blocks)  # page=None must not break chunking


def test_parse_pptx_slides_and_notes():
    blocks = parse_document(_pptx_bytes(), "application/vnd.openxmlformats-officedocument.presentationml.presentation", "deck.pptx")
    assert len(blocks) == 1  # the blank slide is skipped
    assert blocks[0]["page"] == 1 and blocks[0]["section_title"] == "Slide 1: Roadmap"
    assert "Ship parsing support" in blocks[0]["content"] and "Mention the pilot customers" in blocks[0]["content"]


def test_corrupt_office_files_raise_clear_errors():
    import pytest

    with pytest.raises(ValueError, match="Word document"):
        parse_document(b"not a zip", "application/octet-stream", "bad.docx")
    with pytest.raises(ValueError, match="PowerPoint"):
        parse_document(b"not a zip", "application/octet-stream", "bad.pptx")


def test_parse_xlsx_sheets():
    import io

    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Region", "Q1"])
    ws.append(["EMEA", 42])
    ws.append([None, None])  # empty rows are dropped
    wb.create_sheet("Empty")  # empty sheets are dropped
    ws3 = wb.create_sheet("Costs")
    ws3.append(["Item", "Cost"])
    out = io.BytesIO()
    wb.save(out)

    blocks = parse_document(out.getvalue(), "application/octet-stream", "book.xlsx")
    assert [b["section_title"] for b in blocks] == ["Sheet: Sales", "Sheet: Costs"]
    assert blocks[0]["content"] == "Region | Q1\nEMEA | 42"
    assert [b["page"] for b in blocks] == [1, 3]


def test_corrupt_xlsx_raises_clear_error():
    import pytest

    with pytest.raises(ValueError, match="Excel"):
        parse_document(b"not a zip", "application/octet-stream", "bad.xlsx")


def test_pptx_parsing_leaves_the_environment_alone():
    import os

    from amber.ingest import parser

    parser._markitdown.cache_clear()  # force the lazy import to run inside this test
    before = dict(os.environ)
    parse_document(_pptx_bytes(), "application/octet-stream", "deck.pptx")
    assert dict(os.environ) == before
